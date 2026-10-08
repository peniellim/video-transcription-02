#!/usr/bin/env python3
"""
Video Transcription Tool (up to 60 min)
- Auto-detects spoken language and frame rate (incl. 23.976 / 29.97 DF / 59.94 DF)
- Extracts dialogue with FPS-based timecodes (HH:MM:SS:FF, ';' for drop-frame)
- Extracts on-screen text (OCR) with timecode ranges
- Translates source language to English (dialogue + on-screen text)

Requirements:
  ffmpeg + ffprobe on PATH
  pip install faster-whisper easyocr opencv-python openpyxl argostranslate

Usage:
  python video_transcriber.py video.mp4
  python video_transcriber.py video.mp4 --model large-v3 --ocr-interval 1.0 --device cuda
  python video_transcriber.py video.mp4 --no-ocr
"""
import argparse, json, subprocess, sys, tempfile, os
from fractions import Fraction

MAX_SECONDS = 60 * 60 + 5  # 60 min (+ small tolerance)

def preload_cuda12():
    """Colab may ship CUDA 13; ctranslate2 needs the CUDA 12 cuBLAS/cuDNN pip wheels."""
    import glob, ctypes, site
    for sp in site.getsitepackages():
        for pat in ("nvidia/cublas/lib/libcublas.so.12*", "nvidia/cublas/lib/libcublasLt.so.12*",
                    "nvidia/cudnn/lib/libcudnn.so.9*"):
            for f in sorted(glob.glob(os.path.join(sp, pat))):
                try:
                    ctypes.CDLL(f, mode=ctypes.RTLD_GLOBAL)
                except OSError:
                    pass

# ---------------------------------------------------------------- probing
def probe(path):
    out = subprocess.check_output([
        "ffprobe", "-v", "error", "-select_streams", "v:0",
        "-show_entries", "stream=avg_frame_rate,r_frame_rate:format=duration",
        "-of", "json", path])
    d = json.loads(out)
    s = d["streams"][0]
    avg, r = Fraction(s["avg_frame_rate"]), Fraction(s["r_frame_rate"])
    fps = float(avg if avg > 0 else r)
    vfr = avg > 0 and abs(float(avg) - float(r)) > 0.5
    return fps, float(d["format"]["duration"]), vfr

# ---------------------------------------------------------------- timecode
def is_drop(fps):
    return abs(fps - 29.97) < 0.01 or abs(fps - 59.94) < 0.01

def frames_to_tc(frames, fps):
    nominal = round(fps)
    sep = ":"
    if is_drop(fps):
        sep = ";"
        d = 2 if nominal == 30 else 4
        per10 = round(fps * 600)
        D, M = divmod(frames, per10)
        frames += 9 * d * D + (d * ((M - d) // (nominal * 60 - d)) if M > d else 0)
    ff = frames % nominal
    s = frames // nominal
    return f"{s // 3600:02d}:{(s // 60) % 60:02d}:{s % 60:02d}{sep}{ff:02d}"

def sec_to_tc(sec, fps):
    return frames_to_tc(int(round(sec * fps)), fps)

def srt_time(sec):
    ms = int(round(sec * 1000))
    return f"{ms // 3600000:02d}:{ms // 60000 % 60:02d}:{ms // 1000 % 60:02d},{ms % 1000:03d}"

# ---------------------------------------------------------------- translation (text -> English)
def make_translator(src):
    if src == "en":
        return lambda t: t
    try:
        import argostranslate.package as pkg, argostranslate.translate as tr
        if not any(l.code == src for l in tr.get_installed_languages()):
            pkg.update_package_index()
            p = next(p for p in pkg.get_available_packages()
                     if p.from_code == src and p.to_code == "en")
            pkg.install_from_path(p.download())
        return lambda t: tr.translate(t, src, "en")
    except Exception as e:
        print(f"  [warn] text translator unavailable for '{src}': {e}")
        return None

# ---------------------------------------------------------------- dialogue
def transcribe(path, model_name, device, tmp, dur=0.0, progress=lambda *a: None):
    from faster_whisper import WhisperModel
    wav = os.path.join(tmp, "audio.wav")
    subprocess.check_call(["ffmpeg", "-y", "-v", "error", "-i", path, "-vn", "-ac", "1",
                           "-ar", "16000", wav])
    import wave, numpy as np
    with wave.open(wav, "rb") as w:  # load audio ourselves (avoids PyAV version issues)
        audio = np.frombuffer(w.readframes(w.getnframes()), dtype=np.int16).astype(np.float32) / 32768.0
    model = WhisperModel(model_name, device=device,
                         compute_type="float16" if device == "cuda" else "int8")
    print("Transcribing (auto language detection)...")
    segs, info = model.transcribe(audio, language=None, vad_filter=True, beam_size=5)
    src = []
    for sg in segs:
        src.append((sg.start, sg.end, sg.text.strip()))
        progress("Transcribing audio", 0.03 + 0.42 * min(1, sg.end / max(dur, 1)))
    lang = info.language
    print(f"  Detected language: {lang} (p={info.language_probability:.2f})")
    trans = []
    if lang != "en":
        print("Translating dialogue to English...")
        segs2, _ = model.transcribe(audio, language=lang, task="translate",
                                    vad_filter=True, beam_size=5)
        trans = []
        for sg in segs2:
            trans.append((sg.start, sg.end, sg.text.strip()))
            progress("Translating dialogue", 0.45 + 0.13 * min(1, sg.end / max(dur, 1)))
    return lang, src, trans

def align(src, trans):
    """Assign each English segment to the source segment it overlaps most."""
    out = [[] for _ in src]
    for ts, te, tt in trans:
        best, bo = None, 0
        for i, (ss, se, _) in enumerate(src):
            o = min(te, se) - max(ts, ss)
            if o > bo:
                best, bo = i, o
        if best is not None:
            out[best].append(tt)
    return [" ".join(x) for x in out]

# ---------------------------------------------------------------- on-screen text
OCR_LANG = {"ko": "ko", "ja": "ja", "zh": "ch_sim"}

def ocr_video(path, fps, lang, interval, device, total=1, progress=lambda *a: None):
    import cv2, easyocr
    langs = [OCR_LANG.get(lang, lang), "en"] if lang != "en" else ["en"]
    try:
        reader = easyocr.Reader(langs, gpu=(device == "cuda"))
    except Exception:
        reader = easyocr.Reader(["en"], gpu=(device == "cuda"))
    cap = cv2.VideoCapture(path)
    step = max(1, int(round(interval * fps)))
    groups, cur, last_small, last_text, idx = [], None, None, "", 0
    print("Extracting on-screen text...")
    while cap.grab():
        if idx % step == 0:
            progress("Reading on-screen text", 0.60 + 0.34 * min(1, idx / max(total, 1)))
            ok, frame = cap.retrieve()
            if not ok:
                break
            small = cv2.resize(cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY), (64, 36))
            if last_small is None or cv2.absdiff(small, last_small).mean() > 3:
                res = [(b[0][1], b[0][0], t) for b, t, c in reader.readtext(frame) if c > 0.5 and len(t.strip()) > 1]
                last_text = " | ".join(t for _, _, t in sorted(res))
                last_small = small
            if last_text:
                if cur and cur["text"] == last_text:
                    cur["end"] = idx
                else:
                    cur = {"start": idx, "end": idx, "text": last_text}
                    groups.append(cur)
            else:
                cur = None
        if idx % (step * 300) == 0:
            print(f"  {idx / fps / 60:.1f} min scanned")
        idx += 1
    cap.release()
    for g in groups:
        g["end"] += step  # hold until next sample
    return groups

# ---------------------------------------------------------------- main
def process(video, model="large-v3", device="auto", ocr_interval=1.0, do_ocr=True,
            progress=lambda *a: None):
    class a:  # simple namespace
        pass
    a.video, a.model, a.device, a.ocr_interval, a.no_ocr = video, model, device, ocr_interval, not do_ocr
    progress("Analyzing video", 0.01)
    fps, dur, vfr = probe(a.video)
    print(f"FPS: {fps:.3f}{' (drop-frame TC)' if is_drop(fps) else ''} | Duration: {dur / 60:.1f} min")
    if vfr:
        print("  [warn] variable frame rate detected; timecodes use average FPS")
    if dur > MAX_SECONDS:
        raise ValueError("Video exceeds 60 minutes.")
    device = a.device
    if device == "auto":
        try:
            import torch
            device = "cuda" if torch.cuda.is_available() else "cpu"
        except ImportError:
            device = "cpu"

    with tempfile.TemporaryDirectory() as tmp:
        if device == "cuda":
            preload_cuda12()
        try:
            lang, src, trans = transcribe(a.video, a.model, device, tmp, dur, progress)
        except RuntimeError as e:
            if device != "cuda":
                raise
            print(f"  [warn] GPU transcription failed ({e}); falling back to CPU (slower)")
            device = "cpu"
            lang, src, trans = transcribe(a.video, a.model, device, tmp, dur, progress)
    english = align(src, trans) if trans else [t for _, _, t in src]

    ocr = []
    if not a.no_ocr:
        ocr = ocr_video(a.video, fps, lang, a.ocr_interval, device, dur * fps, progress)
        progress("Translating on-screen text", 0.95)
        tl = make_translator(lang)
        for g in ocr:
            g["en"] = tl(g["text"]) if tl else ""

    progress("Writing Excel", 0.97)
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Dialogue"
    ws.append(["#", "TC In", "TC Out", f"Source ({lang})", "English"])
    for i, ((s, e, t), en) in enumerate(zip(src, english), 1):
        ws.append([i, sec_to_tc(s, fps), sec_to_tc(e, fps), t, en])
    ws2 = wb.create_sheet("On-Screen Text")
    ws2.append(["#", "TC In", "TC Out", f"Source ({lang})", "English"])
    for i, g in enumerate(ocr, 1):
        ws2.append([i, frames_to_tc(g["start"], fps), frames_to_tc(g["end"], fps), g["text"], g.get("en", "")])
    ws3 = wb.create_sheet("Info")
    for row in [("File", os.path.basename(a.video)), ("Detected language", lang), ("FPS", round(fps, 3)),
                ("Drop-frame", is_drop(fps)), ("Duration (min)", round(dur / 60, 2))]:
        ws3.append(row)
    for sh in (ws, ws2):
        for c in sh[1]:
            c.font = Font(bold=True)
        for col, w in zip("ABCDE", (6, 14, 14, 60, 60)):
            sh.column_dimensions[col].width = w
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    ws3.column_dimensions["A"].width = 20
    ws3.column_dimensions["B"].width = 40

    base = os.path.splitext(a.video)[0]
    wb.save(base + "_transcript.xlsx")
    with open(base + "_en.srt", "w", encoding="utf-8") as f:
        for i, ((st, en_, _), en) in enumerate(zip(src, english), 1):
            f.write(f"{i}\n{srt_time(st)} --> {srt_time(en_)}\n{en}\n\n")
    print(f"Done: {len(src)} dialogue lines, {len(ocr)} on-screen entries")
    progress("Done", 1.0)
    return base + "_transcript.xlsx", base + "_en.srt", {
        "language": lang, "fps": round(fps, 3), "duration_min": round(dur / 60, 1),
        "dialogue": len(src), "onscreen": len(ocr)}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("video")
    ap.add_argument("--model", default="large-v3")
    ap.add_argument("--device", default="auto", choices=["auto", "cpu", "cuda"])
    ap.add_argument("--ocr-interval", type=float, default=1.0, help="seconds between OCR samples")
    ap.add_argument("--no-ocr", action="store_true")
    a = ap.parse_args()
    process(a.video, a.model, a.device, a.ocr_interval, not a.no_ocr)


if __name__ == "__main__":
    main()
